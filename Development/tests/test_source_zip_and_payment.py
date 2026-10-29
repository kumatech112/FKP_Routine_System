"""Encrypted inner ZIP intake and Payment attachment presentation regressions."""
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from io import BytesIO
import tempfile,unittest
import pyzipper
from openpyxl import Workbook,load_workbook
from routine.files import extract_payload
from routine.settings import default_settings,runtime_settings,task_config
from routine.service import RoutineJob,safe_trace
from routine.config import engine_settings
from routine.payment import prepare_payment
from routine.delivery import package_files
from routine.vendor.fkp.domain import ValidationError


class InnerZipTests(unittest.TestCase):
    def create(self,path):
        with pyzipper.AESZipFile(path,'w',compression=pyzipper.ZIP_DEFLATED) as archive:
            archive.setpassword(b'fixture-source');archive.setencryption(pyzipper.WZ_AES,nbits=256)
            archive.writestr('nested/case.pdf',b'%PDF-1.4 fixture encrypted payload')

    def test_saved_password_extracts_aes_inner_zip_and_keeps_pdf_bytes(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'payload.zip';self.create(source)
            values=default_settings();values['secrets']['zip_password']='fixture-source'
            decrypt=runtime_settings(values,'enforcement')['decrypt']
            work=root/'Process';work.mkdir()
            folder=extract_payload(source,SimpleNamespace(folder=work,check=lambda:None),decrypt,allow_pdf=True)
            self.assertEqual((folder/'nested/case.pdf').read_bytes(),b'%PDF-1.4 fixture encrypted payload')
            self.assertTrue(source.exists())

    def test_missing_or_wrong_source_password_blocks_before_pdf_processing(self):
        for password in ['', 'incorrect']:
            with self.subTest(password_set=bool(password)),tempfile.TemporaryDirectory() as temp:
                root=Path(temp);source=root/'payload.zip';self.create(source)
                work=root/'Process';work.mkdir()
                with self.assertRaisesRegex(ValidationError,'ZIP'):
                    extract_payload(source,SimpleNamespace(folder=work,check=lambda:None),{'archive_password':password},allow_pdf=True)
                self.assertTrue(source.exists())

    def test_archive_password_is_redacted_from_errors(self):
        detail=safe_trace(RuntimeError('failed with fixture-source'),{'decrypt':{'archive_password':'fixture-source'}})
        self.assertNotIn('fixture-source',detail)
        self.assertIn('<redacted>',detail)


class PaymentPresentationTests(unittest.TestCase):
    def test_prepared_and_zipped_payment_hides_query_preserves_values_and_has_no_processed_suffix(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'A1015_PaymentPostWO_Online_R2_29102026.xlsx'
            book=Workbook();book.active.append(['acc']);book.active.append(['000123']);book.save(source);book.close()
            database=root/'fixture.mdb';database.write_bytes(b'fixture MDB')
            values=default_settings();values.update(logs_root=str(root/'Logs'),output_root=str(root/'Completed'))
            settings=engine_settings(task_config(values,'payment_online'));job=RoutineJob(root/'job',settings)
            data=SimpleNamespace(rows=[['000123']],import_rows=[['000123']],headers=['acc'],round_label='R2')
            headers=['No','TcccNo','TYPE_JOB1','จ่าย/ไม่จ่ายวิชาชีพ','FORECAST','collectorname1']
            rows=[[1,'00001234','fixture-job','fixture-election','fixture-forecast','fixture-collector']]
            dao=SimpleNamespace(QueryDefs=[],Close=lambda:None)
            with patch('routine.payment.legacy.extract_source_workbook_data',return_value=data), \
                 patch('routine.payment.legacy.import_rows_into_access'), \
                 patch('routine.payment.legacy.fetch_query_rows',return_value=(headers,rows)), \
                 patch('routine.payment.legacy.apply_round_updates'), \
                 patch('routine.payment.legacy.base.open_dao_database',return_value=(object(),dao)):
                result=prepare_payment(source,database,job,run_actions=False)
            result.verify()
            self.assertEqual(result.output.name,source.name)
            self.assertNotIn('_processed',result.output.name)
            book=load_workbook(result.output)
            self.assertEqual(book['Table1'].sheet_state,'visible')
            self.assertEqual(book.active.title,'Table1')
            self.assertEqual(book['Query1'].sheet_state,'hidden')
            self.assertEqual(list(book['Query1'].values),[tuple(headers),tuple(rows[0])])
            self.assertEqual(book['Table1']['A2'].value,'000123')
            self.assertEqual(book['Table1']['B2'].value,'00001234')
            book.close()
            delivery=root/'Delivery';delivery.mkdir()
            attachment=package_files([result.output],{'zip_attachment':True,'zip_password':'fixture-send'},SimpleNamespace(folder=delivery),result.output.stem)
            self.assertNotIn('_processed',attachment.name)
            with pyzipper.AESZipFile(attachment) as archive:
                archive.setpassword(b'fixture-send')
                self.assertEqual(archive.namelist(),[source.name])
                sent=load_workbook(BytesIO(archive.read(source.name)))
                self.assertEqual(sent['Query1'].sheet_state,'hidden')
                self.assertEqual(sent.active.title,'Table1');sent.close()
            self.assertEqual(database.read_bytes(),b'fixture MDB')


if __name__=='__main__':unittest.main()
