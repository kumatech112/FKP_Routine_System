from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import tempfile,unittest
from PIL import Image
from openpyxl import load_workbook
from routine import ocr
from routine.enforcement import write_workbook
from routine.enforcement_flow import copy_pdf
from routine.pdf_timestamp import parse_timestamp,pdf_timestamp,timestamp_consensus
from routine.vendor.fkp.domain import ValidationError

class TimestampTests(unittest.TestCase):
    def test_red_frame_format_and_calendar_validation(self):
        self.assertEqual(parse_timestamp('พ. 28/10/2026 17:53'),'พ. 28/10/2026 17:53')
        self.assertEqual(parse_timestamp('พ. ๒๘/๑๐/๒๕๖๙ ๑๗:๕๓'),'พ. 28/10/2569 17:53')
        for text in ('31/02/2026 17:53','28/10/2026 24:53','28102026','28/10/2026'):
            self.assertEqual(parse_timestamp(text),'')
    def test_ambiguous_dates_are_never_guessed(self):
        self.assertEqual(parse_timestamp('28/10/2026 17:53 29/10/2026 17:53'),'')
    def test_majority_resolves_similar_digits_and_disagreement_stays_blank(self):
        self.assertEqual(timestamp_consensus(['28/10/2026 17:53','28/10/2026 17:33','6 17:53']),'พ. 28/10/2026 17:53')
        self.assertEqual(timestamp_consensus(['27/10/2026 08:32','27/10/2026 09:32','27/10/2026 09:32']),'อ. 27/10/2026 09:32')
        self.assertEqual(timestamp_consensus(['28/10/2026 17:53','28/10/2026 17:33','6 17:23']),'')
    def test_header_region_excludes_table_dates(self):
        from PIL import ImageDraw
        image=Image.new('RGB',(1000,1000),'white')
        ImageDraw.Draw(image).rectangle((600,300,800,310),fill='black')
        with patch('routine.pdf_timestamp.recognize',return_value=('28/10/2026 17:53',95)) as reader:
            self.assertEqual(pdf_timestamp(image,[{'y':[500,600]}],'exe'),'พ. 28/10/2026 17:53')
            self.assertLess(reader.call_args.args[0].height,500)

class WorkbookTests(unittest.TestCase):
    def test_date_is_final_column_per_record_including_continued_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'result.xlsx'
            review={'version':3,'output':str(path),'tables':[{'id':'t1','group_id':'g','headers':['acc','ชื่อ']},
                {'id':'t2','group_id':'g','headers':['acc','ชื่อ']},{'id':'t3','headers':['acc','ชื่อ']}],
                'cells':[{'table_id':'t1','group_id':'g','record_id':'a','row':1,'column':0,'value':'00123','pdf_datetime':'พ. 28/10/2026 17:53'},
                {'table_id':'t2','group_id':'g','record_id':'a','row':1,'column':1,'value':'ชื่อ นามสกุล','pdf_datetime':'พ. 28/10/2026 17:53'},
                {'table_id':'t1','group_id':'g','record_id':'b','row':2,'column':0,'value':'00456','pdf_datetime':'พฤ. 29/10/2026 09:00'},
                {'table_id':'t3','row':1,'column':0,'value':'00789','pdf_datetime':''}]}
            write_workbook(review)
            book=load_workbook(path)
            self.assertEqual(book.sheetnames,['ตาราง','ตาราง 2'])
            self.assertEqual(list(book.active.values),[('acc','ชื่อ','วันที่อีเมลใน PDF'),
                ('00123','ชื่อ นามสกุล','พ. 28/10/2026 17:53'),('00456',None,'พฤ. 29/10/2026 09:00')])
            self.assertIsNone(book.worksheets[1]['C2'].value)
            self.assertEqual(book.active['C2'].data_type,'s')
            self.assertEqual(book.active.auto_filter.ref,'A1:C3')
            review['cells'][1]['pdf_datetime']='พฤ. 29/10/2026 09:00'
            with self.assertRaises(ValidationError):write_workbook(review)

class BackgroundTests(unittest.TestCase):
    def test_native_txt_spacing_no_console_and_language_check_cached(self):
        with tempfile.TemporaryDirectory() as temp:
            exe=Path(temp)/'tesseract.exe';exe.write_bytes(b'fixture')
            calls=[]
            def invoke(command,**kwargs):
                calls.append((command,kwargs))
                self.assertTrue(kwargs['creationflags'] & 0x08000000)
                self.assertTrue(kwargs['startupinfo'].dwFlags & 1)
                self.assertEqual(kwargs['startupinfo'].wShowWindow,0)
                if '--list-langs' in command:return SimpleNamespace(returncode=0,stdout=b'tha\neng\n')
                base=Path(command[2])
                base.with_suffix('.txt').write_text('ชื่อ นามสกุล\nไม่พบ payment ต่อเนื่อง',encoding='utf-8')
                base.with_suffix('.tsv').write_text('conf\ttext\n90\tน\n85\tา\n',encoding='utf-8')
                return SimpleNamespace(returncode=0,stdout=b'')
            ocr._check_languages.cache_clear()
            with patch('routine.ocr.subprocess.run',side_effect=invoke):
                for _ in range(2):
                    text,score=ocr.recognize(Image.new('RGB',(60,20),'white'),exe)
                    self.assertEqual(text,'ชื่อ นามสกุล\nไม่พบ payment ต่อเนื่อง')
                    self.assertEqual(score,87.5)
            self.assertEqual(sum('--list-langs' in cmd for cmd,kw in calls),1)
            self.assertFalse(any('--version' in cmd for cmd,kw in calls))

class PdfNamesTests(unittest.TestCase):
    def test_pdf_root_has_no_artificial_input_folders_and_collisions_retain_both(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'source.pdf';source.write_bytes(b'PDF one')
            first=copy_pdf(source,root/'PDF',Path('case.pdf'))
            source.write_bytes(b'PDF two')
            second=copy_pdf(source,root/'PDF',Path('case.pdf'))
            self.assertEqual(first.relative_to(root).as_posix(),'PDF/case.pdf')
            self.assertEqual(second.relative_to(root).as_posix(),'PDF/case (2).pdf')
            self.assertEqual(first.read_bytes(),b'PDF one')
            self.assertEqual(second.read_bytes(),b'PDF two')

if __name__=='__main__':unittest.main()
