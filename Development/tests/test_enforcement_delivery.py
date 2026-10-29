"""Enforcement intake and ZIP delivery faults; all mail/Office adapters are isolated."""
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import json
import tempfile
import threading
import unittest

import pyzipper
from routine.enforcement import FIELDS, write_workbook
from routine.matching import enforcement_match
from routine.outlook import Outlook, MailRef
from routine.settings import default_settings, runtime_settings
from routine.vendor.fkp.domain import ValidationError
from routine.vendor.fkp.persistence import BANGKOK, digest, dump
from routine.workflows import FileWorkflowService, finalize_enforcement


class IntakeTests(unittest.TestCase):
    def test_decorated_titles_and_minor_typo_match(self):
        for title in ('แจ้งผลขออนุมัติยึดอสังหาริมทรัพย์_28102026',
            'FW: แจ้งผลขออนุมัติยึดอสังหาริมทรัพย์_28102026 (ทดสอบ)',
            '[TEST] ขออนุมัติยึดอสังหาริมทรัพย์ เพิ่มเติม',
            'แจ้งผลขออนุมัติยึดอสังหาริมทรัพย_28102026'):
            with self.subTest(title=title):
                self.assertIsNotNone(enforcement_match(title,''))

    def test_body_can_classify_when_title_is_unrelated(self):
        found=enforcement_match('ทดสอบระบบ','A1015L_Approval_Enforcement_28102026.zip.pgp')
        self.assertEqual(found['matched_by'],'body_filename')
        self.assertIsNotNone(enforcement_match('เอกสารแนบ','แจ้งผลขออนุมัติยึดอสังหาริมทรัพย์'))

    def test_unrelated_mail_and_wrong_oa_are_not_classified(self):
        for title,body in [('แจ้งผลชำระเงิน','เอกสารทั่วไป'),
                           ('เอกสาร','A1099_Approval_Enforcement_28102026.zip.pgp')]:
            self.assertIsNone(enforcement_match(title,body))

    def test_actual_scanner_uses_body_and_preserves_unread_cutoff(self):
        now=datetime.now(BANGKOK)
        def mail(id,subject,body='',unread=True,age=0):
            return SimpleNamespace(Class=43,EntryID=id,Subject=subject,Body=body,
                UnRead=unread,ReceivedTime=now-timedelta(hours=age))
        mails=[mail('body','ทดสอบ','A1015_Approval_Enforcement_28102026.zip.pgp'),
               mail('decorated','RE: แจ้งผลขออนุมัติยึดอสังหาริมทรัพย์_28102026 (TEST)'),
               mail('read','ขออนุมัติยึดอสังหาริมทรัพย์',unread=False),
               mail('other','ทั่วไป'),mail('old','ขออนุมัติยึดอสังหาริมทรัพย์',age=48)]
        items=SimpleNamespace(Count=len(mails),Sort=lambda *a:None,Item=lambda i:mails[i-1])
        folder=SimpleNamespace(Items=items)
        store=SimpleNamespace(DisplayName='fixture',StoreID='store',GetRootFolder=lambda:folder,
                              GetDefaultFolder=lambda _:folder)
        outlook=Outlook.__new__(Outlook)
        outlook.session=SimpleNamespace(Stores=[store])
        settings=runtime_settings(default_settings(),'enforcement')['scan']
        settings['store_name']='fixture'
        settings['unread_only']=True
        self.assertEqual({r.entry_id for r in outlook.scan(settings)},{'body','decorated'})
        settings.update(processor='payment_online',subject_regex=r'Job_(?P<mail_date>\d{8})')
        mails[:]=[mail('exact','Job_28102026'),mail('extra','FW: Job_28102026 (TEST)')]
        items.Count=len(mails)
        self.assertEqual([r.entry_id for r in outlook.scan(settings)],['exact'])


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)
        self.values=default_settings()
        self.values.update(logs_root=str(self.root/'Logs'),output_root=str(self.root/'Completed'))
        self.values['archive']['archive_processed_files']=False
        self.values['mail']['workflow_recipients']['enforcement']={'to':['fixture-enforcement'],'cc':['fixture-review']}
        self.values['mail'].update(zip_attachment=True,zip_password_required=True,send_immediately=True)
        self.values['secrets']['zip_password']='fixture-only-zip'
        self.source=self.root/'A1015L_Approval_Enforcement_28102026.pdf.pgp'
        self.source.write_bytes(b'fixture encrypted bytes')
        self.reference=MailRef('entry','store','FW: enforcement (TEST)',self.source.name,{})
        self.events=[]
        self.fail_send=False
        test=self
        class FakeOutlook:
            def prepare_forward(self,reference,attachment,settings):
                test.assertEqual(settings['to'],['fixture-enforcement'])
                test.assertEqual(settings['cc'],['fixture-review'])
                with pyzipper.AESZipFile(attachment) as archive:
                    archive.setpassword(b'fixture-only-zip')
                    names=archive.namelist()
                    test.assertIn('Enforcement.xlsx',names)
                    pdfs=[x for x in names if x.endswith('.pdf')]
                    test.assertTrue(pdfs)
                    test.assertFalse(any(x.endswith('.txt') for x in names))
                    for name in pdfs:test.assertEqual(archive.read(name),b'fixture pdf')
                    from io import BytesIO
                    from openpyxl import load_workbook
                    book=load_workbook(BytesIO(archive.read('Enforcement.xlsx')))
                    test.assertEqual(book.sheetnames,['ตาราง'])
                    test.assertEqual(book.active.max_column,21)
                    test.assertEqual(book.active.max_row,2)
                test.events.append('prepared')
                return SimpleNamespace(Close=lambda _:test.events.append('closed'))
            def deliver(self,forwarded,settings):
                test.events.append('mail')
                if test.fail_send:raise RuntimeError('ambiguous delivery')
                return 'sent' if settings['send_immediately'] else 'draft_opened'
        self.outlook=FakeOutlook()

    def tearDown(self):
        self.temp.cleanup()

    def extract(self,source,job,settings,allow_pdf=False):
        self.assertTrue(allow_pdf)
        self.events.append('extract')
        extracted=job.folder/'Extracted'
        (extracted/'one').mkdir(parents=True)
        (extracted/'two').mkdir()
        (extracted/'one/case.pdf').write_bytes(b'fixture pdf')
        (extracted/'two/notes.txt').write_bytes(b'not part of table output')
        return extracted

    def ocr(self,folder,job,executable):
        self.events.append('ocr')
        review={'version':2,'status':'needs_review','page_count':1,'cell_count':21,
            'output':str(job.folder/'Enforcement.xlsx'),'pages':[],
            'tables':[{'id':'t1','headers':list(FIELDS),'rows':1}],
            'cells':[{'id':str(i),'table_id':'t1','row':1,'column':i,'field':field,
                'value':'','confidence':0,'state':'needs_review','image':''} for i,field in enumerate(FIELDS)]}
        dump(job.folder/'enforcement_review.json',review)
        write_workbook(review)
        return review

    def execute(self,ocr=None,sources=None):
        self.service=FileWorkflowService(self.values,'enforcement',threading.Event(),outlook=self.outlook)
        runtime=runtime_settings(self.values,'enforcement')
        with patch('routine.enforcement_flow.extract_payload',side_effect=self.extract),              patch('routine.enforcement_flow.prepare_enforcement',side_effect=ocr or self.ocr):
            return self.service.process_unit(self.reference,sources or self.source,runtime,self.outlook,'unit')

    def state(self):
        return next(iter(self.service.ledger.data['files'].values()))

    def resume(self):
        return finalize_enforcement(Path(self.state()['review']),{},values=self.values,outlook=self.outlook)

    def test_ocr_then_password_zip_then_mail_without_manual_review(self):
        receipt=self.execute()
        self.assertEqual(self.events,['extract','ocr','prepared','mail'])
        self.assertEqual(receipt['status'],'success')
        self.assertEqual({Path(p).suffix for p in receipt['result']['outputs']},{'.xlsx','.zip'})
        review=json.loads(Path(self.state()['review']).read_text(encoding='utf-8'))
        self.assertFalse(review['review_required'])
        self.assertFalse(review['human_reviewed'])
        self.assertTrue(all(c['state']=='needs_review' for c in review['cells']))
        self.assertIsNone(self.execute())
        self.assertEqual(self.events.count('mail'),1)

    def test_missing_recipient_can_be_fixed_without_redoing_ocr(self):
        self.values['mail']['workflow_recipients']['enforcement']['to']=[]
        self.assertEqual(self.execute()['status'],'failed')
        self.assertEqual(self.state()['status'],'ready_to_send')
        self.assertNotIn('mail',self.events)
        self.values['mail']['workflow_recipients']['enforcement']['to']=['fixture-enforcement']
        self.assertEqual(self.resume()['status'],'success')
        self.assertEqual(self.events.count('ocr'),1)

    def test_ocr_failure_never_sends(self):
        def fail(*a):raise RuntimeError('OCR failed')
        self.assertEqual(self.execute(fail)['status'],'failed')
        self.assertNotIn('mail',self.events)
        self.assertTrue(self.source.exists())

    def test_zip_failure_never_sends_and_is_safe_to_resume(self):
        with patch('routine.enforcement_flow.package_files',side_effect=ValidationError('ZIP failed')):
            self.assertEqual(self.execute()['status'],'failed')
        self.assertNotIn('mail',self.events)
        self.assertEqual(self.state()['status'],'ready_to_send')
        self.assertEqual(self.resume()['status'],'success')

    def test_uncertain_delivery_persists_across_restart_and_is_not_sent_again(self):
        self.fail_send=True
        self.assertEqual(self.execute()['status'],'failed')
        self.assertEqual(self.state()['mail'],'unknown')
        self.assertIsNone(self.execute())
        with self.assertRaises(ValidationError):self.resume()
        self.assertEqual(self.events.count('mail'),1)

    def test_draft_keeps_source_and_has_both_pdf_and_excel(self):
        self.values['mail']['send_immediately']=False
        self.values['archive'].update(archive_processed_files=True,processed_archive_folder=str(self.root/'Archived'))
        self.assertEqual(self.execute()['status'],'draft_opened')
        self.assertTrue(self.source.exists())
        self.assertFalse((self.root/'Archived').exists())

    def test_changed_pdf_blocks_resume(self):
        self.values['mail']['workflow_recipients']['enforcement']['to']=[]
        self.execute()
        from routine.vendor.fkp.persistence import read
        context=read(Path(self.state()['review']).parent/'review_context.json')
        Path(context['pdfs'][0]['path']).write_bytes(b'changed')
        self.values['mail']['workflow_recipients']['enforcement']['to']=['fixture-enforcement']
        with self.assertRaisesRegex(ValidationError,'PDF'):self.resume()
        self.assertNotIn('mail',self.events)

    def test_password_is_required_even_if_global_option_is_optional(self):
        self.values['secrets']['zip_password']=''
        self.values['mail']['zip_password_required']=False
        self.assertEqual(self.execute()['status'],'failed')
        self.assertNotIn('mail',self.events)

    def test_multiple_sources_are_one_ocr_and_one_zip(self):
        other=self.root/'A1015L_Reject_Enforcement_28102026.pdf.gpg'
        other.write_bytes(b'another encrypted fixture')
        receipt=self.execute(sources=[self.source,other])
        self.assertEqual(receipt['result']['pdfs'],2)
        self.assertEqual(self.events.count('mail'),1)
        self.assertEqual(self.events.count('ocr'),1)

    def test_old_pre_mail_failure_is_retried_once_through_poll_after_upgrade(self):
        self.values['workflows']['enforcement']['network_folder']=str(self.root)
        self.outlook.scan_workflow=lambda _:[self.reference]
        self.outlook.inline_images=lambda *a:[]
        service=FileWorkflowService(self.values,'enforcement',outlook=self.outlook)
        self.service=service
        def fail(*args):raise RuntimeError('old extraction/OCR failure')
        with patch('routine.enforcement_flow.extract_payload',side_effect=self.extract),patch('routine.enforcement_flow.prepare_enforcement',side_effect=fail):
            self.assertEqual(service.poll()[0]['status'],'failed')
        state=self.state();state.pop('pre_mail_fix');service.ledger.save()
        with patch('routine.enforcement_flow.extract_payload',side_effect=self.extract),patch('routine.enforcement_flow.prepare_enforcement',side_effect=self.ocr):
            self.assertEqual(service.poll()[0]['status'],'success')
            self.assertEqual(service.poll(),[])
        self.assertEqual(self.events.count('mail'),1)
        self.assertEqual(len(self.state()['previous_attempts']),1)

    def test_upgrade_never_retries_pending_unknown_sent_or_draft_outcomes(self):
        from routine.enforcement_flow import can_retry_before_mail
        for outcome in ('pending','unknown','sent','draft_opened'):
            state={'status':'failed','mail':outcome,'collect':'not_started'}
            self.assertFalse(can_retry_before_mail(state))
        self.assertFalse(can_retry_before_mail({'status':'failed','mail':'not_started','collect':'not_started','review':'existing OCR review'}))

    def test_old_unsent_review_can_resume_without_manual_cell_marks(self):
        self.values['mail']['workflow_recipients']['enforcement']['to']=[]
        self.execute()
        self.service.ledger.update('unit',next(iter(self.service.ledger.data['files'])),status='needs_review')
        self.values['mail']['workflow_recipients']['enforcement']['to']=['fixture-enforcement']
        self.assertEqual(self.resume()['status'],'success')


class PayloadAndTableTests(unittest.TestCase):
    def test_direct_pdf_after_pgp_decryption(self):
        from routine.files import extract_payload
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            source=root/'case.pdf.pgp';source.write_bytes(b'encrypted')
            work=root/'work';work.mkdir()
            def decrypt(command,**kwargs):
                Path(command[command.index('--output')+1]).write_bytes(b'%PDF-1.4 fixture bytes')
                return SimpleNamespace(returncode=0)
            with patch('routine.files.subprocess.run',side_effect=decrypt):
                folder=extract_payload(source,SimpleNamespace(folder=work,check=lambda:None),
                    {'enabled':True,'gpg_path':'fixture.exe'},allow_pdf=True)
            self.assertEqual((folder/'case.pdf').read_bytes(),b'%PDF-1.4 fixture bytes')

    def test_approval_reject_bay_and_multiple_pdf_names(self):
        from routine.matching import body_names
        names=['261021_A1015L_A1015L_Reject_Enforcement_21102026.pdf.pgp',
            '261027_A1015L_A1015L_Approval_Enforcement_BAY_27102026_2.pdf.gpg',
            'A1015_Approval_Enforcement_28102026.pgp',
            'A1015_Approval_Enforcement_28102026.zip.pgp']
        self.assertEqual(body_names(' '.join(names),'enforcement'),names)

    def test_workbook_has_only_tables_and_joins_continuation_fragments(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            review={'tables':[{'id':'t','group_id':'g','headers':['acc','ความเห็น'],'rows':1}],
                'cells':[{'id':'h','group_id':'g','column':-2,'header_column':0,'value':'เลขบัญชี'},
                    {'group_id':'g','column':0,'record_id':'row1','row':1,'value':'00123'},
                    {'group_id':'g','column':1,'record_id':'row1','row':1,'value':'ต้นแถว'},
                    {'group_id':'g','column':1,'record_id':'row1','row':1,'value':'ต่อจากหน้าถัดไป'},
                    {'column':-1,'value':'signature must not appear'}],
                'output':str(Path(directory)/'table.xlsx')}
            write_workbook(review)
            book=load_workbook(review['output'])
            self.assertEqual(book.sheetnames,['ตาราง'])
            self.assertEqual(book.active.max_column,2)
            self.assertEqual(book.active.max_row,2)
            self.assertEqual(book.active['A1'].value,'เลขบัญชี')
            self.assertEqual(book.active['A2'].value,'00123')
            self.assertEqual(book.active['B2'].value,'ต้นแถว\nต่อจากหน้าถัดไป')


class BatchIntakeTests(unittest.TestCase):
    def exercise(self,missing=False):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);values=default_settings()
            values.update(logs_root=str(root/'Logs'),output_root=str(root/'Completed'))
            values['workflows']['enforcement']['network_folder']=str(root)
            names=['A1015_Approval_Enforcement_28102026.pdf.pgp','A1015L_Reject_Enforcement_28102026.pdf.gpg']
            for name in names[:1] if missing else names:
                (root/name).write_bytes(name.encode())
            reference=MailRef('entry','store','ทดสอบ',' '.join(names),{})
            adapter=SimpleNamespace(scan_workflow=lambda _: [reference],inline_images=lambda *a:[])
            service=FileWorkflowService(values,'enforcement',outlook=adapter)
            with patch.object(service,'process_unit',return_value={'status':'needs_review'}) as process:
                result=service.poll()
                if missing:
                    process.assert_not_called();self.assertEqual(result,[])
                else:
                    self.assertEqual(process.call_count,1)
                    self.assertEqual({p.name for p in process.call_args.args[1]},set(names))
                    self.assertEqual(len(result),1)

    def test_all_files_in_one_body_are_one_batch(self):
        self.exercise()

    def test_missing_file_waits_for_complete_batch(self):
        self.exercise(True)


if __name__=='__main__':unittest.main()
