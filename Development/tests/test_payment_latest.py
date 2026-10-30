"""Payment latest-payment fields follow Query2 names and survive Excel export."""
from pathlib import Path
import tempfile,unittest
from openpyxl import load_workbook
from routine.vendor import legacy_payment as legacy


class LatestPaymentTests(unittest.TestCase):
    def headers(self):
        # Deliberately different from both Access's order and the Excel output order.
        return ['ยอดชำระล่าสุด','collectorname1','No','FORECAST','วันชำระล่าสุด',
                'TYPE_JOB1','TcccNo','จ่าย/ไม่จ่ายวิชาชีพ']

    def test_named_fields_export_to_t_and_u_without_changing_existing_columns(self):
        headers=self.headers()
        query_rows=[['1,234.50','fixture collector',1,'fixture forecast','29/10/2026',
                     'fixture type','00001234','fixture election']]
        source=['15','date-a','cut','000099','date-b','date-c',100,'time',200,'bucket',
                'type','client','month','reassign']
        values=legacy.build_excel_text_values(headers,query_rows)
        with tempfile.TemporaryDirectory() as temp:
            output=Path(temp)/'payment.xlsx'
            legacy.write_output_workbook(output,import_sheet_name='Table1',query_sheet_name='Query1',
                source_headers=legacy.SOURCE_HEADERS,source_rows=legacy.build_output_rows([source],values),
                query_headers=headers,query_rows=query_rows,overwrite=False)
            book=load_workbook(output)
            try:
                sheet=book['Table1']
                self.assertEqual(sheet.max_column,21)
                self.assertEqual([c.value for c in sheet[1]][-3:],
                                 ['FORECAST','วันชำระล่าสุด','ยอดชำระล่าสุด'])
                self.assertEqual([c.value for c in sheet[2]][:14],source)
                self.assertEqual([c.value for c in sheet[2]][14:19],
                                 ['00001234','fixture collector','fixture type','fixture election','fixture forecast'])
                self.assertEqual(sheet['T2'].value,'29/10/2026')
                self.assertEqual(sheet['U2'].value,'1,234.50')
                self.assertEqual(sheet['T2'].data_type,'s')
                self.assertEqual(sheet['U2'].data_type,'s')
                self.assertEqual(list(book['Query1'].values),[tuple(headers),tuple(query_rows[0])])
                self.assertEqual(book['Query1'].sheet_state,'hidden')
                self.assertEqual(book.active.title,'Table1')
            finally:
                book.close()

    def test_null_is_blank_and_zero_payment_is_retained(self):
        headers=self.headers()
        rows=[[None,'collector',1,'forecast',None,'type','00001','election'],
              [0,'collector',2,'forecast','30/10/2026','type','00002','election']]
        values=legacy.build_excel_text_values(headers,rows)
        self.assertEqual(values[0][-2:],['',''])
        self.assertEqual(values[1][-2:],['30/10/2026','0'])

    def test_missing_new_field_blocks_instead_of_using_an_unrelated_position(self):
        for name in ('วันชำระล่าสุด','ยอดชำระล่าสุด'):
            with self.subTest(field=name):
                headers=[h for h in self.headers() if h!=name]+['unrelated']*40
                with self.assertRaisesRegex(ValueError,name):
                    legacy.build_excel_text_values(headers,[[None]*len(headers)])


if __name__=='__main__':unittest.main()
